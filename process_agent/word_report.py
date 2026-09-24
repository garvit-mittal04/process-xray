"""Build a Word (.docx) version of a Process X-Ray report, in memory."""
import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GREEN = RGBColor(0x0F, 0x6E, 0x56)
GREY = RGBColor(0x5F, 0x5E, 0x5A)
LABELS = {
    "eliminate": "Eliminate", "simplify": "Simplify", "integrate": "Integrate",
    "automate_rules": "Automate (rules)", "automate_ai": "Automate (AI)", "keep_human": "Keep human",
}
VERDICTS = {"go": "GO", "go_with_changes": "GO WITH CHANGES", "rethink": "RETHINK"}
PAYMENT_WORDS = ("payment", "paid", "chase", "remind")


def _shade(cell, hex_fill: str) -> None:
    """Background colour for a table cell (clear shading renders correctly everywhere)."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _table(doc, headers: list[str], rows: list[list], widths_cm: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        _shade(cell, "E1F5EE")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(value))
            run.font.size = Pt(9)
    for row in table.rows:          # widths on every cell so Word and Google Docs agree
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
    doc.add_paragraph()


def _para(doc, text: str = "", bold_prefix: str = "", size: float | None = None,
          color: RGBColor | None = None, style: str | None = None):
    p = doc.add_paragraph(style=style)
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    run = p.add_run(text)
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return p


def build_docx(r) -> bytes:
    pm, an, rp = r.process_map, r.analysis, r.replay
    recs = {x.id: x for x in r.plan.recommendations}
    s = r.plan.settings

    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        doc.styles[name].font.color.rgb = GREEN

    # ---- Cover ----
    doc.add_paragraph("Process X-Ray report", style="Title")
    _para(doc, pm.process_name, size=14, color=GREY)
    _para(doc, f"{r.chat.total_messages} messages analysed, {r.chat.start:%d %b %Y} to "
               f"{r.chat.end:%d %b %Y}. Names are masked (Person 1, 2, 3; Customer A, B...).",
          size=9, color=GREY)

    doc.add_heading("Headline results", level=1)
    _table(doc, ["Cases traced", "Waiting removed", "Chasers avoided", "Time per case"],
           [[rp.cases, f"{rp.waiting_removed_per_week:g} h/week", rp.chasers_avoided,
             f"{rp.avg_cycle_before_h:g} h → {rp.avg_cycle_after_h:g} h"]],
           [4.2, 4.2, 4.2, 4.4])
    _para(doc, an.top_insight, bold_prefix="Biggest insight. ")

    # ---- How the process runs today ----
    doc.add_heading("How the process runs today", level=1)
    _para(doc, pm.summary)
    _para(doc, "; ".join(f"{p} = {role}" for p, role in pm.roles.items()), bold_prefix="Who does what: ")
    _para(doc, f"{an.total_hands_on_hours_per_week} h/week of hands-on work, and "
               f"{an.total_waiting_hours_per_week} h/week spent waiting between steps.",
          bold_prefix="Workload: ")
    if pm.pain_points:
        doc.add_heading("Pain points found in the chat", level=2)
        for p in pm.pain_points:
            _para(doc, p.description, bold_prefix=f"{p.kind.replace('_', ' ').capitalize()}: ",
                  style="List Bullet")

    # ---- Recommendations ----
    doc.add_heading("Recommended automations", level=1)
    _para(doc, "Ranked by value adjusted for risk. Each idea was reviewed by a sceptical "
               "'devil's advocate' agent.", size=9, color=GREY)
    for item in r.ranking:
        rec, c = recs[item.recommendation_id], item.critique
        doc.add_heading(f"#{item.rank} {rec.title}", level=2)
        step_names = {x.id: x.name for x in an.steps}
        covers = ", ".join((lambda n: n[0].lower() + n[1:])(step_names.get(i, i)) for i in rec.step_ids)
        _para(doc, f"{VERDICTS[c.verdict]}, confidence {c.confidence}/5. "
                   f"{LABELS[rec.treatment]}; covers {covers}.", color=GREEN)
        _para(doc, rec.what_it_does)
        _table(doc, ["Work saved", "Waiting removed", "Setup", "Net value"],
               [[f"{rec.hours_saved_per_week:g} h/week", f"{rec.waiting_removed_per_week:g} h/week",
                 f"{rec.setup_days:g} days (Rs {rec.setup_cost_inr:,.0f})",
                 f"Rs {rec.monthly_value_inr:,.0f}/month"]],
               [4.2, 4.2, 4.2, 4.4])
        if rec.monthly_value_inr <= 0:
            _para(doc, "At this volume the tool cost is higher than the staff time saved: the "
                       "benefit is speed and fewer chasers. Prefer free tools.", size=9, color=GREY)
        doc.add_heading("How to build it", level=3)
        for n, step in enumerate(rec.how_it_works, 1):
            _para(doc, step, bold_prefix=f"{n}. ")
        _para(doc, ", ".join(rec.tools), bold_prefix="Tools: ")
        _para(doc, rec.human_in_the_loop, bold_prefix="Human stays in: ")
        if c.risks:
            doc.add_heading("Risks", level=3)
            for risk in c.risks:
                p = _para(doc, f"{risk.risk} ", bold_prefix=f"{risk.severity}/5 ", style="List Bullet")
                fix = p.add_run(f"Fix: {risk.mitigation}")
                fix.italic = True
        _para(doc, c.who_might_resist, bold_prefix="Who may resist: ")
        if c.change_needed:
            _para(doc, c.change_needed, bold_prefix="Change needed: ")

    # ---- Replay ----
    doc.add_heading("Replay: your history with the automations", level=1)
    _para(doc, "The real history was re-run as if the recommended automations had existed. Only "
               "waits an automation directly controls were shortened; customers, transporters and "
               "physical work keep their real timing. Ideas marked RETHINK were left out.")
    ops = [m for m in rp.milestones if not any(w in m.step_name.lower() for w in PAYMENT_WORDS)]
    if ops:
        _table(doc, ["Step reached", "Before (h)", "With automations (h)", "Sooner (h)"],
               [[m.step_name, f"{m.avg_hours_before:g}", f"{m.avg_hours_after:g}",
                 f"{m.avg_hours_before - m.avg_hours_after:.1f}"] for m in ops],
               [8.0, 2.8, 3.4, 2.8])
    if rp.moments:
        doc.add_heading("Moments that would have gone differently", level=2)
        for mo in rp.moments:
            _para(doc, f"{mo.step_name} at {mo.after:%a %d %b %H:%M} instead of "
                       f"{mo.before:%a %d %b %H:%M} ({mo.hours_sooner:g} working hours sooner)",
                  bold_prefix=f"{mo.case_label}: ", style="List Bullet")

    # ---- Step detail ----
    doc.add_heading("Step-by-step analysis", level=1)
    _table(doc, ["Step", "Who", "Per week", "Wait (h)", "Treatment", "Potential"],
           [[f"{x.id}. {x.name}", x.actor, f"{x.frequency_per_week:g}", f"{x.wait_before_hours:g}",
             LABELS[x.treatment], f"{x.potential}/5"] for x in an.steps],
           [4.6, 4.4, 1.6, 1.6, 2.8, 2.0])

    # ---- Method ----
    doc.add_heading("Assumptions and method", level=1)
    for line in [
        f"Staff time valued at Rs {s.staff_cost_per_hour_inr:g}/hour; owner time at Rs {s.owner_cost_per_hour_inr:g}/hour.",
        f"Building an automation costs Rs {s.setup_cost_per_day_inr:,.0f}/day; each chasing message costs about {s.minutes_per_chaser:g} minutes.",
        "Frequencies and waits are measured from real timestamps in working hours (Mon-Sat, 9:00-19:00). Minutes per step are estimates.",
        "AI outputs pass through code guardrails: waits are measured, steps that depend on customers or transporters cannot be sped up, and names are masked before any AI call.",
    ]:
        _para(doc, line, style="List Bullet")

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "Generated by Process X-Ray · estimates to discuss, not guarantees"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = GREY

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
